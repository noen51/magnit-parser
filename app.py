from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
import traceback
from collections import Counter
from datetime import datetime
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.async_api import BrowserContext, Page, async_playwright


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results"
ERROR_FILE = ROOT / "app_error.txt"
DEBUG_FILE = ROOT / "network_debug.jsonl"

LEGACY_PROFILE = ROOT / "browser_profiles" / "yandex_browser"
CURRENT_PROFILE = ROOT / "browser_profile_yandex"

CORE_TOP_LEVEL_NAMES = {
    "Алкоголь",
    "Готовая еда",
    "Молочный прилавок",
    "Овощи и фрукты",
    "Хлеб и выпечка",
    "Бакалея",
    "Консервы",
    "Птица, мясо",
    "Рыба, морепродукты",
    "Заморозка",
    "Сладости",
    "Снеки",
    "Чай, кофе, какао",
    "Вода и напитки",
    "Для детей",
    "Для животных",
    "Гигиена и уход",
    "Для дома и не только",
}


def resolve_profile() -> Path:
    if LEGACY_PROFILE.exists():
        return LEGACY_PROFILE
    return CURRENT_PROFILE


def find_browser() -> Path | None:
    local = Path(os.environ.get("LOCALAPPDATA", ""))
    pf = Path(os.environ.get("PROGRAMFILES", ""))
    pfx86 = Path(os.environ.get("PROGRAMFILES(X86)", ""))

    candidates = [
        local / "Yandex/YandexBrowser/Application/browser.exe",
        pf / "Yandex/YandexBrowser/Application/browser.exe",
        pfx86 / "Yandex/YandexBrowser/Application/browser.exe",
        pfx86 / "Microsoft/Edge/Application/msedge.exe",
        pf / "Microsoft/Edge/Application/msedge.exe",
        local / "Microsoft/Edge/Application/msedge.exe",
        local / "Google/Chrome/Application/chrome.exe",
        pf / "Google/Chrome/Application/chrome.exe",
        pfx86 / "Google/Chrome/Application/chrome.exe",
    ]
    for path in candidates:
        if path.is_file():
            return path
    return None


def write_debug(record: dict[str, Any]) -> None:
    record = {
        "time": datetime.now().isoformat(timespec="seconds"),
        **record,
    }
    with DEBUG_FILE.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")


def parse_category_url(raw_url: str) -> tuple[str, int, str, str]:
    value = raw_url.strip()
    if not value:
        raise ValueError("Category URL is empty.")
    if not value.startswith(("https://", "http://")):
        value = "https://" + value

    parsed = urlparse(value)
    host = (parsed.hostname or "").lower()
    if host != "magnit.ru" and not host.endswith(".magnit.ru"):
        raise ValueError("Only magnit.ru URLs are allowed.")

    match = re.search(r"/catalog/(\d+)", parsed.path)
    if not match:
        raise ValueError("URL must contain /catalog/NUMBER.")

    query = parse_qs(parsed.query)
    store_code = str((query.get("shopCode") or ["320040"])[0])
    store_type = str((query.get("shopType") or ["express"])[0])
    return value, int(match.group(1)), store_code, store_type


def money(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value) / 100.0, 2)
    except (TypeError, ValueError):
        return None


def clean_promotion_end(value: Any) -> Any:
    """Remove API sentinel/epoch dates while preserving real promotion dates."""
    if value in (None, "", 0, "0"):
        return None

    if isinstance(value, (int, float)):
        if value <= 0:
            return None
        seconds = float(value) / 1000 if float(value) > 10_000_000_000 else float(value)
        try:
            parsed = datetime.utcfromtimestamp(seconds)
        except (OverflowError, OSError, ValueError):
            return value
        if parsed.year <= 1971:
            return None
        return parsed.isoformat(timespec="seconds") + "Z"

    text = str(value).strip()
    if not text:
        return None
    if text.startswith("1970-01-01") or text.startswith("1970-01-02"):
        return None
    return value


def safe_file_name(value: str, limit: int = 90) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    return (value or "file")[:limit]


def category_image_url(node: dict[str, Any]) -> str:
    direct = node.get("image")
    if direct:
        return str(direct)
    for entry in node.get("images") or []:
        if isinstance(entry, dict) and entry.get("url"):
            return str(entry["url"])
    return ""


def iter_category_nodes(
    items: list[dict[str, Any]],
    parents: list[str] | None = None,
):
    parent_names = list(parents or [])
    for node in items:
        name = str(node.get("name") or "")
        path = parent_names + [name]
        yield node, path
        yield from iter_category_nodes(node.get("children") or [], path)


def category_ids_with_ancestors(
    path_map: dict[int, list[dict[str, Any]]],
    selected_ids: set[int],
) -> set[int]:
    result: set[int] = set()
    for category_id in selected_ids:
        path = path_map.get(category_id) or []
        for part in path:
            part_id = part.get("id")
            if isinstance(part_id, int):
                result.add(part_id)
        result.add(category_id)
    return result


def flatten_tree(items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[int, list[dict[str, Any]]]]:
    rows: list[dict[str, Any]] = []
    paths: dict[int, list[dict[str, Any]]] = {}

    def walk(
        nodes: list[dict[str, Any]],
        parents: list[dict[str, Any]],
        parent_id: int | None,
    ) -> None:
        for node in nodes:
            node_id = node.get("id")
            name = str(node.get("name") or "")
            current_node = {
                "id": node_id,
                "name": name,
                "seoCode": node.get("seoCode"),
            }
            current_path = parents + [current_node]
            if isinstance(node_id, int):
                paths[node_id] = current_path

            rows.append(
                {
                    "id": node_id,
                    "name": name,
                    "seoCode": node.get("seoCode"),
                    "parentId": parent_id,
                    "depth": len(parents),
                    "fullPath": " > ".join(part["name"] for part in current_path),
                    "image": category_image_url(node),
                    "localImage": str(node.get("localImage") or ""),
                    "imageSource": str(node.get("imageSource") or ("category" if category_image_url(node) else "none")),
                    "childrenCount": len(node.get("children") or []),
                }
            )
            walk(node.get("children") or [], current_path, node_id)

    walk(items, [], None)
    return rows, paths


def leaf_nodes(node: dict[str, Any]) -> list[dict[str, Any]]:
    children = node.get("children") or []
    if not children:
        return [node]
    result: list[dict[str, Any]] = []
    for child in children:
        result.extend(leaf_nodes(child))
    return result




def _number_text(value: str) -> float | None:
    value = str(value or "").strip().replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", value)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


class _ProductPageTextParser(HTMLParser):
    BLOCK_TAGS = {
        "address", "article", "aside", "blockquote", "br", "dd", "div", "dl", "dt",
        "fieldset", "figcaption", "figure", "footer", "form", "h1", "h2", "h3", "h4",
        "h5", "h6", "header", "hr", "li", "main", "nav", "ol", "p", "pre", "section",
        "table", "tbody", "td", "tfoot", "th", "thead", "tr", "ul",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.images: list[dict[str, str]] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")
        if tag == "img":
            data = {str(k): str(v or "") for k, v in attrs}
            self.images.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "noscript", "svg"}:
            if self._skip_depth:
                self._skip_depth -= 1
            return
        if self._skip_depth:
            return
        if tag in self.BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and data:
            self.parts.append(data)

    def text(self) -> str:
        raw = unescape("".join(self.parts)).replace("\r", "\n")
        lines: list[str] = []
        for line in raw.split("\n"):
            line = re.sub(r"[ \t\xa0]+", " ", line).strip()
            if line:
                lines.append(line)
        return "\n".join(lines)


def _exact_heading_span(text: str, heading: str, start: int = 0) -> tuple[int, int] | None:
    pattern = re.compile(rf"(?m)^\s*{re.escape(heading)}\s*$")
    match = pattern.search(text, start)
    if match:
        return match.span()
    return None


def _section_after_heading(text: str, heading: str, next_headings: list[str]) -> str:
    span = _exact_heading_span(text, heading)
    if not span:
        return ""
    start = span[1]
    end = len(text)
    for candidate in next_headings:
        next_span = _exact_heading_span(text, candidate, start)
        if next_span:
            end = min(end, next_span[0])
    value = text[start:end].strip()
    # Excel has a 32,767 character cell limit.
    return value[:30000]


def _nutrition_from_text(text: str) -> dict[str, float | None]:
    result = {
        "nutrition_kcal": None,
        "nutrition_protein": None,
        "nutrition_fat": None,
        "nutrition_carbs": None,
    }
    heading = re.search(r"(?mi)^\s*Пищевая ценность[^\n]*(?:100|100\s*г)[^\n]*$", text)
    if not heading:
        return result
    segment = text[heading.end(): heading.end() + 1200]
    labels = {
        "nutrition_kcal": ["Ккал", "Калории", "Энергетическая ценность"],
        "nutrition_protein": ["Белки", "Белок"],
        "nutrition_fat": ["Жиры", "Жир"],
        "nutrition_carbs": ["Углеводы", "Углевод"],
    }
    for key, variants in labels.items():
        for label in variants:
            match = re.search(
                rf"(?mi)^\s*{re.escape(label)}\s*$\s*\n\s*([^\n]+)",
                segment,
            )
            if match:
                result[key] = _number_text(match.group(1))
                break
    return result


def _fit_size_from_url(url: str) -> tuple[int, int] | None:
    match = re.search(r"/rs:fit:(\d+):(\d+)/", str(url or ""))
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _image_candidates_from_html(raw_html: str, parser: _ProductPageTextParser) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    seen: set[str] = set()

    def add(url: str, alt: str = "") -> None:
        url = unescape(str(url or "")).replace("\\u002F", "/").replace("\\/", "/")
        url = url.strip().strip('"\'')
        if not url.startswith("https://images-foodtech.magnit.ru/"):
            return
        # srcset entries can have a trailing width descriptor.
        url = url.split(" ", 1)[0].strip()
        if url in seen:
            return
        seen.add(url)
        size = _fit_size_from_url(url)
        candidates.append({"url": url, "alt": alt, "size": size})

    for attrs in parser.images:
        alt = attrs.get("alt", "")
        add(attrs.get("src", ""), alt)
        add(attrs.get("data-src", ""), alt)
        for srcset_key in ("srcset", "data-srcset"):
            for part in attrs.get(srcset_key, "").split(","):
                add(part.strip(), alt)

    # Nuxt SSR often repeats image URLs in serialized data. Capture those too.
    for match in re.finditer(r"https://images-foodtech\.magnit\.ru/[^\"'<>\\s]+", raw_html):
        add(match.group(0))
    return candidates


def _choose_hd_product_image(candidates: list[dict[str, Any]], product_name: str) -> tuple[str, str]:
    if not candidates:
        return "", ""
    name = re.sub(r"\s+", " ", str(product_name or "")).strip().lower()

    def score(item: dict[str, Any]) -> tuple[int, int]:
        url = str(item.get("url") or "")
        alt = re.sub(r"\s+", " ", str(item.get("alt") or "")).strip().lower()
        size = item.get("size") or (0, 0)
        area = int(size[0]) * int(size[1])
        points = 0
        if name and alt and (name in alt or alt in name):
            points += 1_000_000_000
        if "/catalog/" in url or "/pim/goods/" in url:
            points += 100_000_000
        if size and min(size) >= 1200:
            points += 10_000_000
        return points + area, area

    best = max(candidates, key=score)
    url = str(best.get("url") or "")
    size = best.get("size")
    size_text = f"{size[0]}x{size[1]}" if size else ""
    return url, size_text


def build_product_detail_url(product: dict[str, Any], store_code: str, store_type: str) -> str:
    product_id = str(product.get("product_id") or "").strip()
    seo = str(product.get("seo_code") or "").strip()
    slug = product_id + (f"-{seo}" if seo else "")
    return (
        f"https://magnit.ru/product/{quote(slug)}"
        f"?shopCode={quote(str(store_code))}&shopType={quote(str(store_type))}"
    )


def parse_product_detail_html(raw_html: str, product_name: str) -> dict[str, Any]:
    parser = _ProductPageTextParser()
    try:
        parser.feed(raw_html)
    except Exception:
        pass
    text = parser.text()
    nutrition = _nutrition_from_text(text)
    composition = _section_after_heading(
        text,
        "Состав",
        ["Характеристики", "Условия хранения", "Документы", "Отзывы"],
    )
    description = _section_after_heading(
        text,
        "Описание",
        ["Состав", "Характеристики", "Условия хранения", "Документы", "Отзывы"],
    )
    hd_url, hd_size = _choose_hd_product_image(
        _image_candidates_from_html(raw_html, parser), product_name
    )
    return {
        **nutrition,
        "composition": composition,
        "description": description,
        "hd_image": hd_url,
        "hd_image_size": hd_size,
    }


async def enrich_product_details(
    context: BrowserContext,
    products: list[dict[str, Any]],
    store_code: str,
    store_type: str,
    concurrency: int = 4,
) -> dict[str, int]:
    """Fetch public product pages to collect ingredients, nutrition and HD image URLs."""
    total = len(products)
    counters = {"details_ok": 0, "composition": 0, "nutrition": 0, "hd_images": 0, "errors": 0}
    semaphore = asyncio.Semaphore(max(1, concurrency))

    async def one(index: int, product: dict[str, Any]) -> None:
        detail_url = build_product_detail_url(product, store_code, store_type)
        product["product_url"] = detail_url
        async with semaphore:
            raw_html = ""
            last_error: Exception | None = None
            for attempt in range(2):
                try:
                    response = await context.request.get(
                        detail_url,
                        timeout=45_000,
                        headers={"accept": "text/html,application/xhtml+xml"},
                    )
                    if response.ok:
                        raw_html = await response.text()
                        break
                    last_error = RuntimeError(f"HTTP {response.status}")
                except Exception as exc:
                    last_error = exc
                await asyncio.sleep(0.8 * (attempt + 1))

            if not raw_html:
                counters["errors"] += 1
                write_debug({
                    "mode": "product_detail_error",
                    "product_id": product.get("product_id"),
                    "url": detail_url,
                    "error": repr(last_error),
                })
                return

            data = parse_product_detail_html(raw_html, str(product.get("name") or ""))
            for key, value in data.items():
                if value not in (None, ""):
                    product[key] = value
            counters["details_ok"] += 1
            if product.get("composition"):
                counters["composition"] += 1
            if any(product.get(key) is not None for key in (
                "nutrition_kcal", "nutrition_protein", "nutrition_fat", "nutrition_carbs"
            )):
                counters["nutrition"] += 1
            if product.get("hd_image"):
                counters["hd_images"] += 1

        if index % 10 == 0 or index == total:
            print(
                f"Product details: {index}/{total}; composition {counters['composition']}; "
                f"nutrition {counters['nutrition']}; HD photos {counters['hd_images']}"
            )

    # Small batches keep load reasonable while still making detailed collection much faster.
    batch_size = max(1, concurrency)
    for start in range(0, total, batch_size):
        batch = products[start:start + batch_size]
        await asyncio.gather(*(one(start + i + 1, product) for i, product in enumerate(batch)))
        await asyncio.sleep(0.15)
    return counters

def normalize_product(
    item: dict[str, Any],
    category: dict[str, Any],
    category_path: list[dict[str, Any]],
) -> dict[str, Any]:
    promo = item.get("promotion") or {}
    ratings = item.get("ratings") or {}
    gallery = item.get("gallery") or []
    image_urls = [
        str(entry.get("url"))
        for entry in gallery
        if isinstance(entry, dict) and entry.get("url")
    ]
    badges = [
        str(entry.get("text"))
        for entry in (item.get("badges") or [])
        if isinstance(entry, dict) and entry.get("text")
    ]

    path_names = [part["name"] for part in category_path]
    path_ids = [part["id"] for part in category_path if isinstance(part.get("id"), int)]

    category_name = category.get("title") or (path_names[-1] if path_names else "")
    full_path = " > ".join(path_names)
    membership = {
        "category_id": category.get("id"),
        "category_name": category_name,
        "category_path": path_names,
        "category_path_ids": path_ids,
        "full_path": full_path,
    }

    return {
        "category_id": category.get("id"),
        "category_name": category_name,
        "category_path": path_names,
        "category_path_ids": path_ids,
        "product_id": item.get("productId") or item.get("id"),
        "name": item.get("name"),
        "price": money(item.get("price")),
        "old_price": money(promo.get("oldPrice")),
        "discount": promo.get("discountPercent"),
        "promotion": bool(promo.get("isPromotion")),
        "promotion_end": clean_promotion_end(promo.get("endDate")),
        "quantity": item.get("quantity"),
        "rating": ratings.get("rating"),
        "reviews": ratings.get("scoresCount"),
        "comments": ratings.get("commentsCount"),
        "pickup_only": bool(item.get("pickupOnly")),
        "adult": bool(item.get("isForAdults")),
        "weighted": bool((item.get("weighted") or {}).get("isWeighted")),
        "store_code": item.get("storeCode"),
        "seo_code": item.get("seoCode"),
        "badges": badges,
        "image": image_urls[0] if image_urls else "",
        "image_urls": image_urls,
        "hd_image": "",
        "hd_image_size": "",
        "product_url": "",
        "composition": "",
        "nutrition_kcal": None,
        "nutrition_protein": None,
        "nutrition_fat": None,
        "nutrition_carbs": None,
        "description": "",
        "all_category_paths": [full_path] if full_path else [],
        "category_memberships": [membership] if full_path else [],
    }


def merge_product(
    products: dict[str, dict[str, Any]],
    product: dict[str, Any],
) -> None:
    product_id = str(product.get("product_id") or "")
    if not product_id:
        return

    existing = products.get(product_id)
    if existing is None:
        products[product_id] = product
        return

    existing_paths = set(existing.get("all_category_paths") or [])
    existing_paths.update(product.get("all_category_paths") or [])
    existing["all_category_paths"] = sorted(existing_paths)

    memberships: dict[tuple[Any, str], dict[str, Any]] = {}
    for membership in (existing.get("category_memberships") or []) + (product.get("category_memberships") or []):
        if not isinstance(membership, dict):
            continue
        key = (membership.get("category_id"), str(membership.get("full_path") or ""))
        memberships[key] = membership
    existing["category_memberships"] = sorted(
        memberships.values(),
        key=lambda item: (str(item.get("full_path") or ""), str(item.get("category_id") or "")),
    )

    existing_path = existing.get("category_path") or []
    new_path = product.get("category_path") or []
    if len(new_path) > len(existing_path):
        for key in ("category_id", "category_name", "category_path", "category_path_ids"):
            existing[key] = product.get(key)

    # Keep the freshest values from the last successful response.
    for key in (
        "price", "old_price", "discount", "promotion", "promotion_end",
        "quantity", "rating", "reviews", "comments", "pickup_only",
        "adult", "weighted", "store_code", "seo_code", "badges", "image", "image_urls",
        "hd_image", "hd_image_size", "product_url", "composition", "nutrition_kcal",
        "nutrition_protein", "nutrition_fat", "nutrition_carbs", "description",
    ):
        if product.get(key) is not None:
            existing[key] = product.get(key)


async def browser_fetch_json(
    page: Page,
    endpoint: str,
    method: str = "GET",
    body: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    result = await page.evaluate(
        """async ({endpoint, method, body}) => {
            try {
                const options = {
                    method,
                    credentials: 'include',
                    cache: 'no-store',
                    headers: {
                        'accept': 'application/json',
                        'content-type': 'application/json',
                        'x-client-name': 'magnit',
                        'x-device-platform': 'Web',
                        'x-new-magnit': 'true'
                    }
                };
                if (body !== null) {
                    options.body = JSON.stringify(body);
                }
                const response = await fetch(endpoint, options);
                const text = await response.text();
                return {
                    ok: response.ok,
                    status: response.status,
                    contentType: response.headers.get('content-type') || '',
                    text
                };
            } catch (error) {
                return {
                    ok: false,
                    status: 0,
                    contentType: '',
                    text: String(error)
                };
            }
        }""",
        {
            "endpoint": endpoint,
            "method": method,
            "body": body,
        },
    )

    preview = str(result.get("text") or "")
    write_debug(
        {
            "mode": "browser_fetch",
            "endpoint": endpoint,
            "method": method,
            "status": result.get("status"),
            "content_type": result.get("contentType"),
            "preview": preview[:500],
        }
    )

    if not result.get("ok"):
        return None

    try:
        data = json.loads(preview)
    except json.JSONDecodeError:
        return None

    return data if isinstance(data, dict) else None


async def fetch_category_tree(
    page: Page,
    store_code: str,
    store_type: str,
) -> tuple[dict[str, Any], str]:
    endpoint = (
        f"/webgate/v3/categories/store/{quote(store_code)}"
        f"?storetype={quote(store_type)}&catalogtype=3"
    )
    tree = await browser_fetch_json(page, endpoint)

    if tree and isinstance(tree.get("items"), list) and tree.get("items"):
        return tree, "live_api"

    fallback = ROOT / "sample_categories.json"
    if fallback.exists():
        try:
            tree = json.loads(fallback.read_text(encoding="utf-8"))
            if isinstance(tree, dict) and isinstance(tree.get("items"), list):
                print("Live category tree failed; local category-tree fallback is used.")
                return tree, "local_fallback"
        except Exception:
            pass

    return {"items": []}, "missing"


async def fetch_one_category(
    page: Page,
    category_id: int,
    store_code: str,
    store_type: str,
    path_map: dict[int, list[dict[str, Any]]],
    products: dict[str, dict[str, Any]],
) -> tuple[int, int | None, str | None]:
    offset = 0
    limit = 32
    page_number = 0
    added_total = 0
    category_title: str | None = None
    total_count: int | None = None

    while page_number < 200:
        page_number += 1
        body = {
            "categories": [category_id],
            "includeAdultGoods": True,
            "pagination": {"limit": limit, "offset": offset},
            "sort": {"order": "desc", "type": "popularity"},
            "storeCode": store_code,
            "storeType": store_type,
            "catalogType": "3",
        }

        data = await browser_fetch_json(
            page,
            "/webgate/v2/goods/search",
            method="POST",
            body=body,
        )
        if not data:
            raise RuntimeError(f"Category API failed for ID {category_id}.")

        category = data.get("category") or {
            "id": category_id,
            "title": path_map.get(category_id, [{"name": str(category_id)}])[-1]["name"],
        }
        category_title = str(category.get("title") or category_title or "")
        category_path = path_map.get(category_id) or [
            {"id": category_id, "name": category_title or str(category_id), "seoCode": None}
        ]

        items = data.get("items") or []
        before = len(products)
        for item in items:
            if not isinstance(item, dict):
                continue
            merge_product(products, normalize_product(item, category, category_path))
        added = len(products) - before
        added_total += max(added, 0)

        pagination = data.get("pagination") or {}
        has_more = pagination.get("hasMore")
        if isinstance(pagination.get("totalCount"), int):
            total_count = pagination["totalCount"]

        print(
            f"  category {category_id}: page {page_number}, "
            f"received {len(items)}, unique total {len(products)}"
        )

        if has_more is False or not items:
            break
        next_offset = pagination.get("nextOffset")
        if isinstance(next_offset, int):
            offset = next_offset
        else:
            offset += limit

        await page.wait_for_timeout(350)

    return added_total, total_count, category_title


async def download_images(
    context: BrowserContext,
    products: list[dict[str, Any]],
    output_dir: Path,
) -> int:
    images_dir = output_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    downloaded = 0

    for index, product in enumerate(products, start=1):
        url = str(product.get("hd_image") or product.get("image") or "")
        if not url:
            continue
        product_id = str(product.get("product_id") or index)
        name = safe_file_name(str(product.get("name") or "product"), 55)
        target = images_dir / f"{product_id}_{name}.webp"
        product["local_image"] = str(Path("images") / target.name)
        product["downloaded_image_url"] = url

        if target.exists():
            continue
        try:
            response = await context.request.get(url, timeout=60_000)
            if response.ok:
                target.write_bytes(await response.body())
                downloaded += 1
                if downloaded % 20 == 0:
                    print(f"Downloaded images: {downloaded}")
        except Exception as exc:
            write_debug({"mode": "image_error", "url": url, "error": repr(exc)})

    return downloaded


async def download_category_images(
    context: BrowserContext,
    tree: dict[str, Any],
    output_dir: Path,
    include_ids: set[int] | None = None,
    products: list[dict[str, Any]] | None = None,
) -> int:
    images_dir = output_dir / "category_images"
    images_dir.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, Any]] = []
    downloaded = 0

    # If Magnit has no dedicated image for a leaf category, use a real product
    # image from that category as a representative fallback. The manifest and
    # Excel clearly mark the source so it is never confused with an official
    # category banner.
    fallback_by_category: dict[int, str] = {}
    for product in products or []:
        product_image = str(product.get("image") or "")
        if not product_image:
            continue
        for membership in product.get("category_memberships") or []:
            if not isinstance(membership, dict):
                continue
            for category_id in membership.get("category_path_ids") or []:
                if isinstance(category_id, int):
                    fallback_by_category.setdefault(category_id, product_image)

    for node, category_path in iter_category_nodes(tree.get("items") or []):
        category_id = node.get("id")
        if include_ids is not None:
            if not isinstance(category_id, int) or category_id not in include_ids:
                continue

        category_name = str(node.get("name") or "category")
        official_url = category_image_url(node)
        fallback_url = fallback_by_category.get(category_id, "") if isinstance(category_id, int) else ""
        url = official_url or fallback_url
        image_source = "category" if official_url else ("product_fallback" if fallback_url else "none")
        local_path = ""
        status = "no_image"

        if url:
            file_id = str(category_id if category_id is not None else "unknown")
            safe_name = safe_file_name(category_name, 60)
            target = images_dir / f"{file_id}_{safe_name}.webp"
            local_path = str(Path("category_images") / target.name)
            node["localImage"] = local_path
            node["imageSource"] = image_source

            if target.exists() and target.stat().st_size > 0:
                status = "already_exists" if image_source == "category" else "already_exists_fallback_product"
            else:
                try:
                    response = await context.request.get(url, timeout=60_000)
                    if response.ok:
                        target.write_bytes(await response.body())
                        downloaded += 1
                        status = "downloaded" if image_source == "category" else "downloaded_fallback_product"
                        if downloaded % 20 == 0:
                            print(f"Downloaded category images: {downloaded}")
                    else:
                        status = f"http_{response.status}"
                        write_debug(
                            {
                                "mode": "category_image_http_error",
                                "category_id": category_id,
                                "url": url,
                                "image_source": image_source,
                                "status": response.status,
                            }
                        )
                except Exception as exc:
                    status = "error"
                    write_debug(
                        {
                            "mode": "category_image_error",
                            "category_id": category_id,
                            "url": url,
                            "image_source": image_source,
                            "error": repr(exc),
                        }
                    )
        else:
            node["imageSource"] = "none"

        manifest.append(
            {
                "category_id": category_id,
                "category_name": category_name,
                "category_path": " > ".join(category_path),
                "image_url": url,
                "official_category_image_url": official_url,
                "fallback_product_image_url": fallback_url,
                "image_source": image_source,
                "local_image": local_path,
                "status": status,
            }
        )

    (output_dir / "category_images.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Category images downloaded: {downloaded}")
    return downloaded


def unique_sheet_name(base: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", base).strip() or "Category"
    cleaned = cleaned[:31]
    candidate = cleaned
    number = 2
    while candidate in used:
        suffix = f" {number}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        number += 1
    used.add(candidate)
    return candidate


def product_row(product: dict[str, Any]) -> list[Any]:
    path = list(product.get("category_path") or [])
    return [
        path[0] if len(path) > 0 else "",
        path[1] if len(path) > 1 else "",
        path[2] if len(path) > 2 else "",
        path[3] if len(path) > 3 else "",
        " > ".join(path),
        " | ".join(product.get("all_category_paths") or []),
        product.get("category_id"),
        product.get("product_id"),
        product.get("name"),
        "",  # Embedded thumbnail is anchored into this cell later.
        product.get("price"),
        product.get("old_price"),
        product.get("discount"),
        product.get("promotion"),
        product.get("promotion_end"),
        product.get("quantity"),
        product.get("rating"),
        product.get("reviews"),
        product.get("comments"),
        product.get("pickup_only"),
        product.get("adult"),
        product.get("weighted"),
        product.get("store_code"),
        product.get("seo_code"),
        ", ".join(product.get("badges") or []),
        product.get("composition", ""),
        product.get("nutrition_kcal"),
        product.get("nutrition_protein"),
        product.get("nutrition_fat"),
        product.get("nutrition_carbs"),
        product.get("description", ""),
        product.get("product_url", ""),
        product.get("image"),
        product.get("hd_image", ""),
        product.get("hd_image_size", ""),
        product.get("local_image", ""),
    ]


PRODUCT_HEADERS = [
    "Главная категория",
    "Категория",
    "Подкатегория",
    "Уровень 4",
    "Полный путь",
    "Все категории товара",
    "ID категории",
    "ID товара",
    "Название товара",
    "Фото",
    "Цена, ₽",
    "Старая цена, ₽",
    "Скидка, %",
    "Акция",
    "Окончание акции",
    "Остаток",
    "Рейтинг",
    "Оценок",
    "Отзывов",
    "Только самовывоз",
    "18+",
    "Весовой",
    "Код магазина",
    "SEO-код",
    "Метки",
    "Состав",
    "Ккал / 100 г",
    "Белки / 100 г",
    "Жиры / 100 г",
    "Углеводы / 100 г",
    "Описание",
    "Страница товара",
    "Фото из каталога",
    "HD фото",
    "HD размер",
    "Локальное HD фото",
]


PRODUCT_WIDTHS = {
    "Главная категория": 24,
    "Категория": 26,
    "Подкатегория": 26,
    "Уровень 4": 24,
    "Полный путь": 52,
    "Все категории товара": 60,
    "ID категории": 13,
    "ID товара": 17,
    "Название товара": 55,
    "Фото": 15,
    "Цена, ₽": 13,
    "Старая цена, ₽": 15,
    "Скидка, %": 11,
    "Акция": 10,
    "Окончание акции": 22,
    "Остаток": 10,
    "Рейтинг": 10,
    "Оценок": 10,
    "Отзывов": 10,
    "Только самовывоз": 17,
    "18+": 8,
    "Весовой": 10,
    "Код магазина": 14,
    "SEO-код": 38,
    "Метки": 25,
    "Состав": 90,
    "Ккал / 100 г": 15,
    "Белки / 100 г": 15,
    "Жиры / 100 г": 15,
    "Углеводы / 100 г": 18,
    "Описание": 80,
    "Страница товара": 60,
    "Фото из каталога": 60,
    "HD фото": 60,
    "HD размер": 14,
    "Локальное HD фото": 44,
}


def style_table_sheet(sheet, row_count: int) -> None:
    red_fill = PatternFill("solid", fgColor="E21D2D")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="E0E0E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in sheet[1]:
        cell.fill = red_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in sheet.iter_rows(min_row=2, max_row=max(2, row_count + 1)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    sheet.freeze_panes = "A2"
    if row_count > 0:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{row_count + 1}"

    header_index = {str(cell.value): cell.column for cell in sheet[1]}
    for header, width in PRODUCT_WIDTHS.items():
        column = header_index.get(header)
        if column:
            sheet.column_dimensions[get_column_letter(column)].width = width

    price_col = header_index.get("Цена, ₽")
    old_price_col = header_index.get("Старая цена, ₽")
    for row_number in range(2, row_count + 2):
        if price_col:
            sheet.cell(row_number, price_col).number_format = '#,##0.00 "₽"'
        if old_price_col:
            sheet.cell(row_number, old_price_col).number_format = '#,##0.00 "₽"'
        for link_header in ("Страница товара", "Фото из каталога", "HD фото"):
            link_col = header_index.get(link_header)
            if link_col:
                cell = sheet.cell(row_number, link_col)
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.style = "Hyperlink"

    discount_col = header_index.get("Скидка, %")
    if row_count > 0 and discount_col:
        letter = get_column_letter(discount_col)
        sheet.conditional_formatting.add(
            f"{letter}2:{letter}{row_count + 1}",
            ColorScaleRule(
                start_type="min", start_color="FFFFFF",
                mid_type="percentile", mid_value=50, mid_color="FFF2CC",
                end_type="max", end_color="F4CCCC",
            ),
        )


def _resolve_local_product_image(output_dir: Path, product: dict[str, Any]) -> Path | None:
    raw = str(product.get("local_image") or "").strip()
    if not raw:
        return None
    # JSON/Excel uses Windows-style paths, while this also keeps offline tests
    # portable on other operating systems.
    relative = Path(raw.replace("\\", os.sep).replace("/", os.sep))
    candidate = output_dir / relative
    if candidate.is_file() and candidate.stat().st_size > 0:
        return candidate
    return None


def embed_product_images(
    sheet,
    products: list[dict[str, Any]],
    output_dir: Path,
    max_px: int = 76,
) -> int:
    """Embed a small preview into Excel while keeping the downloaded HD file untouched."""
    embedded = 0
    pillow_warning_printed = False

    for row_number, product in enumerate(products, start=2):
        image_path = _resolve_local_product_image(output_dir, product)
        if image_path is None:
            continue

        try:
            image = XLImage(str(image_path))
            width = float(image.width or max_px)
            height = float(image.height or max_px)
            scale = min(max_px / width, max_px / height, 1.0)
            image.width = max(1, int(width * scale))
            image.height = max(1, int(height * scale))
            photo_col = next((cell.column for cell in sheet[1] if cell.value == "Фото"), 10)
            image.anchor = f"{get_column_letter(photo_col)}{row_number}"
            sheet.add_image(image)
            sheet.row_dimensions[row_number].height = 62
            embedded += 1
        except ImportError as exc:
            if not pillow_warning_printed:
                print("WARNING: Pillow is required to embed photos into Excel.")
                print(r"Install it: .venv\Scripts\python.exe -m pip install pillow")
                write_debug({"mode": "excel_image_pillow_missing", "error": repr(exc)})
                pillow_warning_printed = True
            break
        except Exception as exc:
            write_debug(
                {
                    "mode": "excel_image_error",
                    "product_id": product.get("product_id"),
                    "path": str(image_path),
                    "error": repr(exc),
                }
            )

    return embedded


def write_product_sheet(
    sheet,
    products: list[dict[str, Any]],
    output_dir: Path,
    embed_images: bool = False,
) -> int:
    sheet.append(PRODUCT_HEADERS)
    for product in products:
        sheet.append(product_row(product))
    style_table_sheet(sheet, len(products))
    if embed_images:
        return embed_product_images(sheet, products, output_dir)
    return 0


def product_memberships(products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, Any, str]] = set()
    for product in products:
        product_id = str(product.get("product_id") or "")
        for membership in product.get("category_memberships") or []:
            if not isinstance(membership, dict):
                continue
            full_path = str(membership.get("full_path") or "")
            category_id = membership.get("category_id")
            key = (product_id, category_id, full_path)
            if key in seen:
                continue
            seen.add(key)
            path = list(membership.get("category_path") or [])
            rows.append(
                {
                    "product_id": product_id,
                    "product_name": product.get("name"),
                    "category_id": category_id,
                    "category_name": membership.get("category_name"),
                    "category_path": path,
                    "category_path_ids": list(membership.get("category_path_ids") or []),
                    "full_path": full_path,
                }
            )
    return sorted(rows, key=lambda row: (row["full_path"], str(row["product_name"] or ""), row["product_id"]))


def build_excel(
    products: list[dict[str, Any]],
    tree: dict[str, Any],
    summary: dict[str, Any],
    output: Path,
) -> None:
    category_rows, _ = flatten_tree(tree.get("items") or [])
    memberships = product_memberships(products)
    wb = Workbook()
    all_sheet = wb.active
    all_sheet.title = "Все товары"
    # Embed actual downloaded thumbnails only on the main sheet. Category sheets
    # keep links/local paths so the XLSX does not duplicate every image and grow
    # to an unnecessarily huge size.
    embedded_images = write_product_sheet(
        all_sheet, products, output.parent, embed_images=True
    )

    used_names = {"Все товары", "Товар-Категория", "Статистика", "Категории", "Сводка"}
    by_top: dict[str, list[dict[str, Any]]] = {}
    for product in products:
        path = product.get("category_path") or []
        top = str(path[0] if path else "Без категории")
        by_top.setdefault(top, []).append(product)

    for top, rows in sorted(by_top.items()):
        sheet = wb.create_sheet(unique_sheet_name(top, used_names))
        write_product_sheet(sheet, rows, output.parent, embed_images=False)

    relation = wb.create_sheet("Товар-Категория")
    relation.append([
        "ID товара", "Название товара", "ID категории", "Название категории",
        "Главная категория", "Категория", "Подкатегория", "Уровень 4", "Полный путь",
    ])
    for item in memberships:
        path = item.get("category_path") or []
        relation.append([
            item.get("product_id"), item.get("product_name"), item.get("category_id"),
            item.get("category_name"),
            path[0] if len(path) > 0 else "",
            path[1] if len(path) > 1 else "",
            path[2] if len(path) > 2 else "",
            path[3] if len(path) > 3 else "",
            item.get("full_path"),
        ])
    for cell in relation[1]:
        cell.fill = PatternFill("solid", fgColor="E21D2D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    relation.freeze_panes = "A2"
    relation.auto_filter.ref = relation.dimensions
    for index, width in enumerate([18, 55, 14, 34, 25, 30, 34, 28, 80], start=1):
        relation.column_dimensions[get_column_letter(index)].width = width

    stats = wb.create_sheet("Статистика")
    stats.append(["Уровень", "Категория", "Уникальных товаров"])
    top_products: dict[str, set[str]] = {}
    path_products: dict[str, set[str]] = {}
    for item in memberships:
        product_id = str(item.get("product_id") or "")
        path = item.get("category_path") or []
        top = str(path[0] if path else "Без категории")
        full_path = str(item.get("full_path") or "Без категории")
        top_products.setdefault(top, set()).add(product_id)
        path_products.setdefault(full_path, set()).add(product_id)

    for name in sorted(top_products):
        stats.append(["Главная категория", name, len(top_products[name])])
    for name in sorted(path_products):
        stats.append(["Полный путь", name, len(path_products[name])])

    for cell in stats[1]:
        cell.fill = PatternFill("solid", fgColor="E21D2D")
        cell.font = Font(color="FFFFFF", bold=True)
    stats.freeze_panes = "A2"
    stats.auto_filter.ref = stats.dimensions
    stats.column_dimensions["A"].width = 22
    stats.column_dimensions["B"].width = 75
    stats.column_dimensions["C"].width = 24

    category_product_sets: dict[int, set[str]] = {}
    for item in memberships:
        category_id = item.get("category_id")
        if isinstance(category_id, int):
            category_product_sets.setdefault(category_id, set()).add(str(item.get("product_id") or ""))

    cats = wb.create_sheet("Категории")
    cats.append(
        [
            "ID", "Название", "SEO-код", "ID родителя", "Уровень", "Полный путь",
            "Дочерних категорий", "Собрано товаров", "Ссылка на фото категории",
            "Локальное фото категории", "Источник фото",
        ]
    )
    for row in category_rows:
        cats.append(
            [
                row["id"], row["name"], row["seoCode"], row["parentId"], row["depth"],
                row["fullPath"], row["childrenCount"],
                len(category_product_sets.get(row["id"], set())),
                row["image"], row.get("localImage", ""), row.get("imageSource", "none"),
            ]
        )
    for cell in cats[1]:
        cell.fill = PatternFill("solid", fgColor="E21D2D")
        cell.font = Font(color="FFFFFF", bold=True)
    cats.freeze_panes = "A2"
    cats.auto_filter.ref = cats.dimensions
    widths = [14, 34, 40, 14, 10, 70, 18, 18, 60, 46, 22]
    for index, width in enumerate(widths, start=1):
        cats.column_dimensions[get_column_letter(index)].width = width
    for row_number in range(2, cats.max_row + 1):
        image_link_cell = cats.cell(row_number, 9)
        if image_link_cell.value:
            image_link_cell.hyperlink = str(image_link_cell.value)
            image_link_cell.style = "Hyperlink"

    overview = wb.create_sheet("Сводка")
    overview.append(["Показатель", "Значение"])
    enhanced_summary = {
        **summary,
        "product_category_links": len(memberships),
        "excel_embedded_product_images": embedded_images,
    }
    for key, value in enhanced_summary.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        overview.append([key, value])
    for cell in overview[1]:
        cell.fill = PatternFill("solid", fgColor="E21D2D")
        cell.font = Font(color="FFFFFF", bold=True)
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 90

    wb.save(output)


def save_result_files(
    output_dir: Path,
    products: dict[str, dict[str, Any]],
    tree: dict[str, Any],
    summary: dict[str, Any],
    errors: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    product_list = sorted(
        products.values(),
        key=lambda product: (
            " > ".join(product.get("category_path") or []),
            str(product.get("name") or ""),
        ),
    )
    (output_dir / "products.json").write_text(
        json.dumps(product_list, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output_dir / "categories.json").write_text(
        json.dumps(tree, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output_dir / "errors.json").write_text(
        json.dumps(errors, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    build_excel(product_list, tree, summary, output_dir / "products.xlsx")
    return product_list


async def create_browser():
    browser_path = find_browser()
    if browser_path is None:
        raise RuntimeError("Yandex Browser, Edge or Chrome was not found.")

    playwright = await async_playwright().start()
    context = await playwright.chromium.launch_persistent_context(
        user_data_dir=str(resolve_profile()),
        executable_path=str(browser_path),
        headless=False,
        no_viewport=True,
        args=["--start-maximized"],
        locale="ru-RU",
        service_workers="allow",
    )
    page = context.pages[0] if context.pages else await context.new_page()
    page.set_default_navigation_timeout(120_000)
    return playwright, context, page


async def ensure_page_ready(page: Page, url: str) -> None:
    await page.goto(url, wait_until="domcontentloaded")
    await page.wait_for_timeout(3500)


async def parse_single_category(
    url: str,
    download_product_photos: bool,
    download_category_photos: bool,
) -> Path:
    url, category_id, store_code, store_type = parse_category_url(url)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = RESULTS / stamp
    output_dir.mkdir(parents=True, exist_ok=True)

    DEBUG_FILE.write_text("", encoding="utf-8")
    errors: list[dict[str, Any]] = []
    products: dict[str, dict[str, Any]] = {}

    playwright, context, page = await create_browser()
    try:
        print(f"Using profile: {resolve_profile()}")
        print(f"Store: {store_code} ({store_type})")
        await ensure_page_ready(page, url)

        tree, tree_source = await fetch_category_tree(page, store_code, store_type)
        _, path_map = flatten_tree(tree.get("items") or [])
        category_image_ids = category_ids_with_ancestors(path_map, {category_id})

        try:
            await fetch_one_category(
                page,
                category_id,
                store_code,
                store_type,
                path_map,
                products,
            )
        except Exception as first_error:
            print("Automatic API call failed.")
            print("In the browser select pickup/store and open the category.")
            input("Then return here and press Enter... ")
            await page.wait_for_timeout(1500)
            try:
                await fetch_one_category(
                    page,
                    category_id,
                    store_code,
                    store_type,
                    path_map,
                    products,
                )
            except Exception as second_error:
                errors.append(
                    {
                        "category_id": category_id,
                        "first_error": repr(first_error),
                        "second_error": repr(second_error),
                    }
                )

        if not products:
            raise RuntimeError("No products were received.")

        product_list = list(products.values())
        print("Collecting composition, nutrition and HD image URLs from product pages...")
        detail_stats = await enrich_product_details(
            context, product_list, store_code, store_type
        )
        downloaded_product_images = 0
        if download_product_photos:
            downloaded_product_images = await download_images(
                context, product_list, output_dir
            )

        downloaded_category_images = 0
        if download_category_photos:
            downloaded_category_images = await download_category_images(
                context, tree, output_dir, category_image_ids, product_list
            )

        summary = {
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "mode": "single_category",
            "source_url": url,
            "store_code": store_code,
            "store_type": store_type,
            "category_id": category_id,
            "products_count": len(products),
            "product_category_links": len(product_memberships(list(products.values()))),
            "category_tree_source": tree_source,
            "product_details": detail_stats,
            "downloaded_product_images": downloaded_product_images,
            "downloaded_category_images": downloaded_category_images,
            "errors_count": len(errors),
        }
        save_result_files(output_dir, products, tree, summary, errors)
        return output_dir
    finally:
        await context.close()
        await playwright.stop()


async def get_tree_for_store(
    page: Page,
    store_code: str,
    store_type: str,
) -> tuple[dict[str, Any], str, dict[int, list[dict[str, Any]]]]:
    main_url = f"https://magnit.ru/?shopCode={quote(store_code)}&shopType={quote(store_type)}"
    await ensure_page_ready(page, main_url)
    tree, tree_source = await fetch_category_tree(page, store_code, store_type)
    if not tree.get("items"):
        raise RuntimeError("Category tree was not received.")
    _, path_map = flatten_tree(tree.get("items") or [])
    return tree, tree_source, path_map


async def parse_categories_batch(
    selected_categories: list[dict[str, Any]],
    store_code: str,
    store_type: str,
    mode: str,
    download_product_photos: bool,
    download_category_photos: bool,
) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = RESULTS / stamp
    output_dir.mkdir(parents=True, exist_ok=True)
    DEBUG_FILE.write_text("", encoding="utf-8")

    products: dict[str, dict[str, Any]] = {}
    errors: list[dict[str, Any]] = []

    playwright, context, page = await create_browser()
    try:
        tree, tree_source, path_map = await get_tree_for_store(page, store_code, store_type)
        total = len(selected_categories)
        selected_category_ids = {
            int(category["id"])
            for category in selected_categories
            if isinstance(category.get("id"), int)
        }
        category_image_ids = category_ids_with_ancestors(
            path_map, selected_category_ids
        )

        for index, category in enumerate(selected_categories, start=1):
            category_id = category.get("id")
            category_name = category.get("name")
            if not isinstance(category_id, int):
                continue
            print(f"[{index}/{total}] {category_name} (ID {category_id})")
            try:
                await fetch_one_category(
                    page,
                    category_id,
                    store_code,
                    store_type,
                    path_map,
                    products,
                )
            except Exception as exc:
                errors.append(
                    {
                        "category_id": category_id,
                        "category_name": category_name,
                        "error": repr(exc),
                    }
                )
                print(f"  ERROR: {exc}")

            # Save progress after every category.
            progress = {
                "mode": mode,
                "processed": index,
                "total": total,
                "products_count": len(products),
                "errors_count": len(errors),
                "updated_at": datetime.now().isoformat(timespec="seconds"),
            }
            (output_dir / "progress.json").write_text(
                json.dumps(progress, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if index % 5 == 0 or index == total:
                partial_summary = {
                    **progress,
                    "store_code": store_code,
                    "store_type": store_type,
                    "category_tree_source": tree_source,
                }
                save_result_files(output_dir, products, tree, partial_summary, errors)

            await page.wait_for_timeout(450)

        product_list = list(products.values())
        print("Collecting composition, nutrition and HD image URLs from product pages...")
        detail_stats = await enrich_product_details(
            context, product_list, store_code, store_type
        )
        downloaded_product_images = 0
        if download_product_photos:
            downloaded_product_images = await download_images(
                context, product_list, output_dir
            )

        downloaded_category_images = 0
        if download_category_photos:
            downloaded_category_images = await download_category_images(
                context, tree, output_dir, category_image_ids, product_list
            )

        summary = {
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "mode": mode,
            "store_code": store_code,
            "store_type": store_type,
            "categories_processed": total,
            "products_count": len(products),
            "product_category_links": len(product_memberships(list(products.values()))),
            "category_tree_source": tree_source,
            "product_details": detail_stats,
            "downloaded_product_images": downloaded_product_images,
            "downloaded_category_images": downloaded_category_images,
            "errors_count": len(errors),
        }
        save_result_files(output_dir, products, tree, summary, errors)
        return output_dir
    finally:
        await context.close()
        await playwright.stop()


async def choose_section_and_parse() -> Path:
    store_code = input("Store code [320040]: ").strip() or "320040"
    store_type = input("Store type [express]: ").strip() or "express"
    download_product_photos = (
        input("Download HD product photos + embed previews into Excel? y/N: ").strip().lower() == "y"
    )
    download_category_photos = (
        input("Download category photos? Y/n: ").strip().lower() != "n"
    )

    playwright, context, page = await create_browser()
    try:
        tree, _, _ = await get_tree_for_store(page, store_code, store_type)
        top_level = tree.get("items") or []
        print()
        for index, node in enumerate(top_level, start=1):
            print(f"{index:>2} - {node.get('name')}")

        selection = input("Select one section number: ").strip()
        selected_index = int(selection) - 1
        if selected_index < 0 or selected_index >= len(top_level):
            raise ValueError("Wrong section number.")
        selected_node = top_level[selected_index]
        categories = leaf_nodes(selected_node)
        print(
            f"Selected: {selected_node.get('name')}; "
            f"leaf categories: {len(categories)}"
        )
    finally:
        await context.close()
        await playwright.stop()

    return await parse_categories_batch(
        categories,
        store_code,
        store_type,
        mode=f"section:{selected_node.get('name')}",
        download_product_photos=download_product_photos,
        download_category_photos=download_category_photos,
    )


async def parse_core_catalog() -> Path:
    store_code = input("Store code [320040]: ").strip() or "320040"
    store_type = input("Store type [express]: ").strip() or "express"
    download_product_photos = (
        input("Download all HD product photos + embed previews into Excel? y/N: ").strip().lower() == "y"
    )
    download_category_photos = (
        input("Download category photos? Y/n: ").strip().lower() != "n"
    )

    playwright, context, page = await create_browser()
    try:
        tree, _, _ = await get_tree_for_store(page, store_code, store_type)
        selected_top = [
            node
            for node in (tree.get("items") or [])
            if str(node.get("name") or "") in CORE_TOP_LEVEL_NAMES
        ]
        categories: list[dict[str, Any]] = []
        for node in selected_top:
            categories.extend(leaf_nodes(node))

        print("Core sections:")
        for node in selected_top:
            print(f" - {node.get('name')}")
        print(f"Leaf categories to parse: {len(categories)}")
        confirmation = input("This can take a long time. Type YES to continue: ").strip()
        if confirmation != "YES":
            raise RuntimeError("Cancelled.")
    finally:
        await context.close()
        await playwright.stop()

    return await parse_categories_batch(
        categories,
        store_code,
        store_type,
        mode="core_catalog",
        download_product_photos=download_product_photos,
        download_category_photos=download_category_photos,
    )


def offline_test() -> Path:
    goods_path = ROOT / "sample_goods.json"
    tree_path = ROOT / "sample_categories.json"
    if not goods_path.exists() or not tree_path.exists():
        raise RuntimeError("sample_goods.json or sample_categories.json is missing.")

    goods = json.loads(goods_path.read_text(encoding="utf-8"))
    tree = json.loads(tree_path.read_text(encoding="utf-8"))
    _, path_map = flatten_tree(tree.get("items") or [])

    products: dict[str, dict[str, Any]] = {}
    category = goods.get("category") or {}
    category_id = int(category.get("id") or 0)
    path = path_map.get(category_id) or [
        {"id": category_id, "name": str(category.get("title") or category_id), "seoCode": None}
    ]
    for item in goods.get("items") or []:
        merge_product(products, normalize_product(item, category, path))

    output_dir = RESULTS / datetime.now().strftime("%Y%m%d_%H%M%S_TEST_V7")
    output_dir.mkdir(parents=True, exist_ok=True)
    summary = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "mode": "offline_test_v7",
        "products_count": len(products),
        "expected_category_path": " > ".join(part["name"] for part in path),
    }
    save_result_files(output_dir, products, tree, summary, [])
    return output_dir


def print_menu() -> None:
    print()
    print("=" * 68)
    print("MAGNIT PARSER V7")
    print("=" * 68)
    print("1 - Parse one category")
    print("2 - Parse one main section with subcategories")
    print("3 - Parse full core catalog (long)")
    print("4 - Offline Excel/category test")
    print("0 - Exit")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--test", action="store_true")
    args = parser.parse_args()

    if args.test:
        result = offline_test()
        print("TEST V7 OK")
        print(result)
        return 0

    print_menu()
    choice = input("> ").strip()

    if choice == "0":
        return 0
    if choice == "4":
        result = offline_test()
    elif choice == "1":
        url = input("Paste full Magnit category URL:\n> ").strip()
        product_photos = (
            input("Download HD product photos + embed previews into Excel? y/N: ").strip().lower() == "y"
        )
        category_photos = (
            input("Download category photos? Y/n: ").strip().lower() != "n"
        )
        result = asyncio.run(
            parse_single_category(url, product_photos, category_photos)
        )
    elif choice == "2":
        result = asyncio.run(choose_section_and_parse())
    elif choice == "3":
        result = asyncio.run(parse_core_catalog())
    else:
        print("Unknown option.")
        return 1

    print()
    print("DONE")
    print(result)
    print("Main file: products.xlsx")
    print("Embedded product previews: sheet 'Все товары', column 'Фото' (HD originals stay in images)")
    print("HD product images folder: images")
    print("Category images folder: category_images")
    print("Category image manifest: category_images.json")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception:
        ERROR_FILE.write_text(traceback.format_exc(), encoding="utf-8")
        print(traceback.format_exc())
        raise SystemExit(1)
